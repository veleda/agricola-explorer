"""
Import a tournament from an Excel workbook in the NM tournament template format
(sheet "Oppsett": starting list + one block per round + optional final).

Reads:
* tournament name and dates from the title rows (A1/A2)
* participants from the starting list (col B short name, col K full name)
* seating per game from each "ROUND n" block (Table / StartNo / Name columns)
* scores from the "ScoreRn" column
* final round ("ROUND n (Final)") and its point weight

Works both with sheets saved by Excel (cached formula values) and with files
exported from Agricola Explorer (no cached values): names are then resolved via
the start number lookup or by following simple cell references like =X172.
"""

from __future__ import annotations

import io
import re
from collections import OrderedDict

import openpyxl
from openpyxl.utils import column_index_from_string

ROUND_RE = re.compile(r"^\s*ROUND\s+(\d+)", re.I)
REF_RE = re.compile(r"^=\$?([A-Z]{1,3})\$?(\d+)$")


def _clean(v):
    if isinstance(v, str):
        v = " ".join(v.split())
        return v or None
    return v


def _is_error(v):
    return isinstance(v, str) and v.startswith("#")


def parse_tournament_xlsx(data: bytes, filename: str = "") -> dict:
    try:
        wbv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        wbf = openpyxl.load_workbook(io.BytesIO(data))
    except Exception as e:
        raise ValueError(f"Could not read the Excel file ({e.__class__.__name__}). Is it an .xlsx file?")
    sheet = "Oppsett" if "Oppsett" in wbv.sheetnames else wbv.sheetnames[0]
    wv, wf = wbv[sheet], wbf[sheet]
    warnings: list[str] = []
    max_row = wv.max_row

    def raw(r, c):
        return _clean(wv.cell(r, c).value)

    # ── starting list ──────────────────────────────────────────────────────
    header_row = None
    for r in range(1, min(max_row, 40) + 1):
        a = raw(r, 1)
        if isinstance(a, str) and a.lower().startswith("initstno"):
            header_row = r
            break
    participants: "OrderedDict[str, dict]" = OrderedDict()
    start_map: dict[int, str] = {}
    if header_row:
        r = header_row + 1
        while r <= max_row and (raw(r, 1) is not None or raw(r, 2) is not None):
            short, full = raw(r, 2), raw(r, 11)
            if isinstance(short, str):
                participants.setdefault(short.casefold(), {"name": short, "fullName": full if isinstance(full, str) else None})
            sn, sname = raw(r, 8), raw(r, 6)
            if isinstance(sn, (int, float)) and isinstance(sname, str):
                start_map[int(sn)] = sname
            r += 1
    else:
        warnings.append("No starting list found (header 'InitStNo1') – participants are taken from the rounds.")

    def resolve_name(r, c, start_col=None, depth=0):
        v = raw(r, c)
        if isinstance(v, str) and not _is_error(v):
            return v
        f = wf.cell(r, c).value
        if isinstance(f, str) and depth < 4:
            m = REF_RE.match(f.replace(" ", ""))
            if m:
                return resolve_name(int(m.group(2)), column_index_from_string(m.group(1)), None, depth + 1)
        if start_col:
            sn = raw(r, start_col)
            if isinstance(sn, (int, float)):
                return start_map.get(int(sn))
        return None

    # ── round blocks ───────────────────────────────────────────────────────
    title_rows = []
    for r in range(1, max_row + 1):
        a = raw(r, 1)
        if isinstance(a, str) and ROUND_RE.match(a):
            title_rows.append(r)
    if not title_rows:
        raise ValueError("No 'ROUND 1', 'ROUND 2' … blocks found. Is this the tournament template (sheet 'Oppsett')?")

    blocks = []
    for bi, tr in enumerate(title_rows):
        text = raw(tr, 1)
        is_final = "final" in text.lower()
        end = title_rows[bi + 1] - 1 if bi + 1 < len(title_rows) else max_row
        hdr = next((r for r in range(tr + 1, min(tr + 4, end) + 1) if str(raw(r, 1) or "").lower() == "table"), None)
        if not hdr:
            warnings.append(f"{text}: header row with 'Table' not found – skipped.")
            continue
        cols = {}
        for c in range(1, min(wv.max_column, 40) + 1):
            h = raw(hdr, c)
            if not isinstance(h, str):
                continue
            hl = h.lower()
            if hl == "name" and "name" not in cols:
                cols["name"] = c
            elif hl == "startno" and "start" not in cols:
                cols["start"] = c
            elif hl.startswith("score") and "score" not in cols:
                cols["score"] = c
        if "name" not in cols:
            warnings.append(f"{text}: no 'Name' column – skipped.")
            continue
        tables: "OrderedDict[int, list[str]]" = OrderedDict()
        scores, empty_run = [], 0
        for r in range(hdr + 1, end + 1):
            a = raw(r, 1)
            if isinstance(a, str) and (a.upper().startswith("FINAL STANDINGS") or ROUND_RE.match(a)):
                break
            name = resolve_name(r, cols["name"], cols.get("start"))
            if not name or name == 0:
                empty_run += 1
                if empty_run >= 2:
                    break
                continue
            empty_run = 0
            if not isinstance(a, (int, float)):
                continue  # listed but not seated (did not play)
            t = int(a)
            tables.setdefault(t, []).append(name)
            sc = raw(r, cols["score"]) if "score" in cols else None
            if isinstance(sc, (int, float)):
                scores.append({"name": name, "table": t, "total": int(round(sc))})
        weight = raw(tr, 6) if is_final else None
        blocks.append({"title": text, "final": is_final, "tables": [tables[k] for k in sorted(tables)],
                       "scores": scores, "weight": weight if isinstance(weight, (int, float)) else None})

    prelim = [b for b in blocks if not b["final"]]
    final = next((b for b in blocks if b["final"]), None)
    if not prelim:
        raise ValueError("No preliminary rounds with players found.")

    # participants from rounds that are not in the starting list
    for b in blocks:
        for tbl in b["tables"]:
            for n in tbl:
                if n.casefold() not in participants:
                    participants[n.casefold()] = {"name": n, "fullName": None}
                    if header_row:
                        warnings.append(f"'{n}' is seated in {b['title']} but not on the starting list – added.")

    # tournament name + dates from the title rows
    a1, a2 = raw(1, 1), raw(2, 1)
    name = None
    if isinstance(a1, str):
        m = re.search(r"setup for\s+(.+?)(?:,\s*\d+\s*pl\b|$)", a1, re.I)
        name = (m.group(1) if m else a1).strip(" ,")
    if not name:
        name = re.sub(r"\.xlsx?$", "", filename or "Imported tournament", flags=re.I)
    start_date = end_date = ""
    if isinstance(a2, str):
        m = re.search(r"(\d{4})\D{1,3}(\d{1,2})\D{1,3}(\d{1,2})(?:\s*[-–]\s*(\d{1,2}))?", a2)
        if m:
            y, mo, d1, d2 = int(m.group(1)), int(m.group(2)), int(m.group(3)), m.group(4)
            if 1 <= mo <= 12 and 1 <= d1 <= 31:
                start_date = f"{y:04d}-{mo:02d}-{d1:02d}"
                end_date = f"{y:04d}-{mo:02d}-{int(d2):02d}" if d2 else start_date

    final_seated = bool(final and any(final["tables"]))
    sizes = [len(t) for t in prelim[0]["tables"]]
    all_sizes = [len(t) for b in prelim for t in b["tables"]]
    score_list = []
    for gi, b in enumerate(prelim, start=1):
        score_list += [{**s, "game": str(gi)} for s in b["scores"]]
    if final:
        score_list += [{**s, "game": "Final"} for s in final["scores"]]
    if final and not final_seated:
        warnings.append("The final is not seated yet in the sheet – use 'Seat final by standings' when the preliminary games are done.")

    return {
        "name": name, "startDate": start_date, "endDate": end_date,
        "participants": list(participants.values()),
        "rounds": [b["tables"] for b in prelim],
        "final": final["tables"] if final_seated else None,
        "hasFinal": bool(final),
        "finalWeight": (final or {}).get("weight") or 1.25,
        "numGames": len(prelim),
        "numTables": len(sizes),
        "sizes": sizes,
        "tableFill": "5" if any(s == 5 for s in all_sizes) else "3",
        "scores": score_list,
        "warnings": warnings,
        "sheet": sheet,
    }
