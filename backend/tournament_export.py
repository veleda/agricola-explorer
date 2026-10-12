"""
Tournament Excel export that follows the NM tournament spreadsheet template
(sheet "Oppsett" with starting list, one block per round, optional final and
final standings, plus a "Printout" table overview).

Layout, formulas, colours, borders and number formats are copied from the
template file in backend/templates/tournament_template.xlsx (a cleaned copy of
the organisers' sheet with all names and values removed). The sheet adapts to
any number of players, tables and rounds.

The template has a few manual "copy → paste values → sort" steps (table ranks
after tie adjustment, overall rank, alphabetical and ranked lists). Those cells
are filled in here with values computed exactly as the template's formulas
would compute them, so the exported file is complete and consistent.
"""

from __future__ import annotations

import io
import os
import random
import re
import unicodedata
from copy import copy
from statistics import mean

import openpyxl
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "templates", "tournament_template.xlsx")

# Template rows used as style sources --------------------------------------
# Starting list: separator line every 4 rows (row 11 = end of a group of four)
START_STYLE = {"header": 7, "first": 8, "mid_first_group": 9, "mid": 12, "group_last": 11, "last": 35}
# per round type: (title, subtitle, header, block_first, group_first, mid, group_last, block_last)
ROUND_STYLE = {
    "r1":     (37, 38, 39, 40, 44, 41, 43, 67),
    "mid":    (70, 71, 72, 73, 77, 74, 76, 100),
    "last":   (169, 170, 171, 172, 176, 173, 175, 199),
    "final":  (202, 203, 204, 205, 209, 206, 208, 232),
}
FINAL_STANDINGS_STYLE = {"title": 235, "note": 236, "header": 237, "first": 238, "mid": 239, "group_last": 241, "last": 265}
NCOLS = 31  # A..AE

TIE_NOTE = "Tie at this table – the template says to adjust manually for ties if your rules require it."


# ── helpers ────────────────────────────────────────────────────────────────

def _nb_key(name: str):
    """Norwegian alphabetical order (A–Z, then Æ, Ø, Å), case-insensitive."""
    s = name.strip().casefold()
    out = []
    for ch in s:
        if ch in "æä":
            out.append("{1")
        elif ch in "øö":
            out.append("{2")
        elif ch == "å":
            out.append("{3")
        else:
            base = unicodedata.normalize("NFKD", ch)
            base = "".join(c for c in base if not unicodedata.combining(c))
            out.append(base)
    return "".join(out)


def _ident(name: str) -> str:
    return " ".join((name or "").split()).casefold()


def _to_int(v):
    if v is None:
        return None
    m = re.search(r"\d+", str(v))
    return int(m.group()) if m else None


FINAL_GAME = 10_000  # sort key for a game registered as "Finale"/"Final"


def _game_no(v):
    """Game number from the score sheet; 'Finale', 'Final', 'F' etc. mark the final round."""
    if v is None:
        return None
    txt = str(v).strip()
    if re.fullmatch(r"(?i)(finale?|f)(\s*\d*)?", txt):
        return FINAL_GAME
    return _to_int(txt)


def _excel_rank(value, values):
    """Excel RANK(value, range) descending: 1 + number of strictly greater values."""
    return 1 + sum(1 for v in values if v > value)


def _points(rank: int) -> int:
    return 10 if rank == 1 else 6 if rank == 2 else 3 if rank == 3 else 1


class _Styler:
    def __init__(self, src_ws, dst_ws):
        self.src, self.dst = src_ws, dst_ws

    def cell(self, src_row, src_col, dst_row, dst_col=None):
        s = self.src.cell(src_row, src_col)
        d = self.dst.cell(dst_row, dst_col or src_col)
        if s.has_style:
            d._style = copy(s._style)
        return d

    def row(self, src_row, dst_row, ncols=NCOLS):
        for c in range(1, ncols + 1):
            self.cell(src_row, c, dst_row)


def _put(ws, row, col, value, comment=None):
    c = ws.cell(row, col)
    c.value = value
    if comment:
        c.comment = Comment(comment, "Agricola Explorer")
    return c


# ── data preparation ───────────────────────────────────────────────────────

def participant_label(p: dict) -> str:
    return " ".join((p.get("displayName") or p.get("name") or p.get("username") or "").split())


def _prepare(tournament: dict, scores: list[dict], include_empty_draw_rounds: bool = False):
    """Return (players, rounds, notes).

    players: list of dicts {key, name, full, init_no, start_no, rand}
    rounds:  list of dicts {game, tables: [(table_label, [player_key, ...])],
                            scores: {key: total}, ids: {key: score_id}}
    Seating comes from the tournament draw when there is one, otherwise from
    the table numbers on the registered scores.
    """
    notes = []
    # Latest score per (game, player)
    latest: dict[tuple[int, str], dict] = {}
    for s in sorted(scores, key=lambda x: x.get("timestamp") or ""):
        g = _game_no(s.get("gameNumber"))
        name = " ".join((s.get("name") or "").split())
        if not name:
            continue
        if g is None:
            notes.append(f"Score for {name} ({s.get('total')}) has no game number and was skipped.")
            continue
        latest[(g, _ident(name))] = {**s, "name": name, "game": g}

    # Players: registered participants first, then anyone else with a score
    display: dict[str, str] = {}
    full_names = {}
    for p in tournament.get("players") or []:
        label = participant_label(p)
        if label:
            display.setdefault(_ident(label), label)
        for cand in (p.get("username"), p.get("displayName"), p.get("name")):
            if cand:
                full_names.setdefault(_ident(cand), label)
    for (_, key), s in latest.items():
        display.setdefault(key, s["name"])
    keys = sorted(display, key=lambda k: _nb_key(display[k]))
    rng = random.Random(str(tournament.get("id") or tournament.get("name")))
    rand = {k: rng.random() for k in keys}
    by_rand = sorted(keys, key=lambda k: rand[k])
    players = []
    for i, k in enumerate(keys, start=1):
        players.append({
            "key": k, "name": display[k], "full": full_names.get(k, display[k]),
            "init_no": i, "rand": rand[k], "start_no": by_rand.index(k) + 1,
        })

    # Draw: game number -> [[names at table 1], [names at table 2], ...]
    draw = tournament.get("draw") or {}
    drawn: dict[int, list[list[str]]] = {i + 1: rnd for i, rnd in enumerate(draw.get("rounds") or [])}
    if draw.get("final"):
        drawn[FINAL_GAME] = draw["final"]

    games = {g for (g, _) in latest}
    if include_empty_draw_rounds:
        games |= set(drawn)
    rounds = []
    for g in sorted(games):
        entries = [s for (gg, _), s in latest.items() if gg == g]
        tables: dict = {}
        seated = set()
        for ti, names in enumerate(drawn.get(g) or [], start=1):
            ks = [_ident(n) for n in names]
            tables[ti] = ks
            seated.update(ks)
        for s in entries:
            k = _ident(s["name"])
            if k in seated:
                continue
            t = _to_int(s.get("tableNumber"))
            tables.setdefault(t if t is not None else "?", []).append(k)
        if "?" in tables:
            notes.append(f"Game {g if g != FINAL_GAME else 'Final'}: {len(tables['?'])} score(s) without table number were grouped as table '?'.")
        order = sorted(tables, key=lambda t: (t == "?", t if t != "?" else 0))
        rounds.append({
            "game": g,
            "tables": [(t, tables[t]) for t in order],
            "scores": {_ident(s["name"]): s.get("total") for s in entries},
            "ids": {_ident(s["name"]): s.get("id") for s in entries},
        })
    return players, rounds, notes


# ── workbook builder ───────────────────────────────────────────────────────

def build_tournament_workbook(tournament: dict, scores: list[dict]) -> bytes:
    players, rounds, notes = _prepare(tournament, scores)
    if not players or not rounds:
        raise ValueError("No scores with game numbers registered for this tournament")

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    src_o, src_p = wb["Oppsett"], wb["Printout"]
    ws = wb.create_sheet("_Oppsett")
    st = _Styler(src_o, ws)

    N = len(players)
    P = {p["key"]: p for p in players}
    start_no = {p["key"]: p["start_no"] for p in players}
    final_by_name = any(r["game"] == FINAL_GAME for r in rounds)
    has_final = (bool(tournament.get("hasFinal")) or final_by_name) and len(rounds) >= 2
    weight = float(tournament.get("finalWeight") or 1.25)
    R = len(rounds)
    n_normal = R - 1 if has_final else R

    # ── Title + starting list ──────────────────────────────────────────────
    first_list_row, last_list_row = 8, 8 + N - 1
    for r in (1, 2, 3):
        st.cell(r, 1, r)
    _put(ws, 1, 1, f"AGRICOLA tournament setup for {tournament.get('name', '')}, {N} pl, "
                   f"{n_normal} rd{' + final' if has_final else ''} ")
    dates = tournament.get("startDate") or ""
    if tournament.get("endDate") and tournament.get("endDate") != dates:
        dates += f" – {tournament['endDate']}"
    _put(ws, 2, 1, dates)
    _put(ws, 3, 1, "Format: table allocations and scores as registered in Agricola Explorer")
    st.cell(5, 1, 5); st.cell(5, 5, 5); st.cell(6, 5, 6)
    _put(ws, 5, 1, "Starting list (sorted by first names col. K) ")
    _put(ws, 5, 5, f"Copy a{first_list_row}:c{last_list_row} to col E and paste values. ")
    _put(ws, 6, 5, f"Sort e{first_list_row}:g{last_list_row} on col G (RandNo), Asc")
    for c in (1, 2, 3, 5, 6, 7, 8, 11):
        st.cell(7, c, 7)
    for c, h in zip((1, 2, 3, 5, 6, 7, 8, 11), ("InitStNo1", "Name", "RandNo", "InitStNo1", "Name", "RandNo", "StartNo", "Players, full names")):
        _put(ws, 7, c, h)
    by_start = sorted(players, key=lambda p: p["start_no"])
    for i in range(N):
        r = first_list_row + i
        src = (START_STYLE["first"] if i == 0 else START_STYLE["last"] if i == N - 1 else
               START_STYLE["group_last"] if i % 4 == 3 else
               START_STYLE["mid_first_group"] if i < 4 else START_STYLE["mid"])
        for c in (1, 2, 3, 5, 6, 7, 8, 11, 12):
            st.cell(src, c, r)
        p = players[i]
        _put(ws, r, 1, p["init_no"]); _put(ws, r, 2, p["name"]); _put(ws, r, 3, "=RAND()")
        q = by_start[i]
        _put(ws, r, 5, q["init_no"]); _put(ws, r, 6, q["name"]); _put(ws, r, 7, round(q["rand"], 15)); _put(ws, r, 8, q["start_no"])
        _put(ws, r, 11, p["full"])
    LIST_H = f"H${first_list_row}:H${last_list_row}"
    LIST_F = f"F${first_list_row}:F${last_list_row}"

    # ── Rounds ─────────────────────────────────────────────────────────────
    title_row = last_list_row + 2
    cum = {k: (0.0, 0.0) for k in P}          # cumulative (ZTurnP, ZRelSc) after previous round
    prev = None                               # info about previous round's lists
    standings_src = None                      # where the final ranked list lives
    final_group_sizes = None                  # final table sizes (for separator lines)

    for idx, rnd in enumerate(rounds):
        k = idx + 1
        is_final = has_final and k == R
        if is_final:
            kind = "final"
        elif k == 1:
            kind = "r1"
        elif k == n_normal:
            kind = "last"
        else:
            kind = "mid"
        sty = ROUND_STYLE[kind]
        ds = title_row + 3
        de = ds + N - 1

        # Build row order: [(table_label, [keys])] + absent group
        played = set()
        groups = []
        for t, keys in rnd["tables"]:
            keys = [x for x in keys if x in P and x not in played and x in rnd["scores"]]
            if not keys:
                continue
            played.update(keys)
            if is_final and prev:
                keys.sort(key=lambda x: prev["rank"][x])
            else:
                keys.sort(key=lambda x: start_no[x])
            groups.append((t, keys, False))
        absent = [x for x in P if x not in played]
        if absent:
            absent.sort(key=lambda x: (prev["rank"][x] if (is_final and prev) else start_no[x]))
            groups.append(("", absent, True))

        if is_final:
            final_group_sizes = [len(g[1]) for g in groups]

        # Python-side computation (mirrors the template formulas)
        row_of, group_range = {}, {}
        r = ds
        for gi, (t, keys, is_abs) in enumerate(groups):
            for x in keys:
                row_of[x] = r
                group_range[x] = (ds + sum(len(g[1]) for g in groups[:gi]), ds + sum(len(g[1]) for g in groups[:gi + 1]) - 1)
                r += 1
        F, G, H, K = {}, {}, {}, {}
        ties = set()
        for t, keys, is_abs in groups:
            if is_abs:
                for x in keys:
                    F[x], G[x], H[x], K[x] = None, 0, None, 0
                continue
            vals = [float(rnd["scores"].get(x) or 0) for x in keys]
            avg = mean(vals)
            for x, v in zip(keys, vals):
                F[x] = rnd["scores"].get(x)
                G[x] = v - avg
                H[x] = _excel_rank(v, vals)
                if vals.count(v) > 1:
                    ties.add(x)
                K[x] = _points(H[x]) * (weight if is_final else 1)
        if k == 1:
            M = {x: K[x] for x in P}
            Nn = {x: G[x] for x in P}
        else:
            M = {x: cum[x][0] + K[x] for x in P}
            Nn = {x: cum[x][1] + G[x] for x in P}

        # Overall rank (P column / ranked list)
        if is_final:
            Pr, offset = {}, 0
            for t, keys, is_abs in groups:
                ordered = sorted(keys, key=lambda x: (-M[x], -Nn[x], _nb_key(P[x]["name"])))
                for i, x in enumerate(ordered):
                    Pr[x] = offset + i + 1
                offset += len(keys)
        else:
            ordered = sorted(P, key=lambda x: (-M[x], -Nn[x], _nb_key(P[x]["name"])))
            Pr = {x: i + 1 for i, x in enumerate(ordered)}
        full_ties = {x for x in P if sum(1 for y in P if M[y] == M[x] and Nn[y] == Nn[x]) > 1}

        # ── header rows ────────────────────────────────────────────────────
        st.row(sty[0], title_row); st.row(sty[1], title_row + 1); st.row(sty[2], title_row + 2)
        if kind == "final":
            _put(ws, title_row, 1, f"ROUND {k} (Final)")
            _put(ws, title_row, 3, f"Weight given to points in R{k} =")
            _put(ws, title_row, 6, weight)
            weight_ref = f"$F${title_row}"
        else:
            _put(ws, title_row, 1, f"ROUND {k}")
        _put(ws, title_row, 10, f"Copy h{ds}:i{de}, paste values and manually adjust for ties")
        if k > 1:
            _put(ws, title_row, 16, f"Copy o{ds}:o{de}, paste values" +
                 (f" and manually resolve ties using ZRelScR{k}" if kind == "final" else f"; resolve ties using ZRelScR{k}"))
        if kind in ("r1", "mid"):
            _put(ws, title_row, 23, f"Copy r{ds}:u{de}, paste values; sort on col. X")
            _put(ws, title_row, 28, f"Copy x{ds}:z{de}, paste values; sort on col. AD and then AE")
        elif kind == "last":
            _put(ws, title_row, 23, f"Copy r{ds}:u{de}, paste values; sort on col. T and then on col. U, and adjust Rank")
        else:
            _put(ws, title_row, 23, f"Copy r{ds}:u{de}, paste values; sort on col. W")

        sub = {1: "Allocation to tables and scores" if k == 1 else "Allocation to tables"}
        if k > 1:
            sub[13] = f"Raw standings after R{k}"
        sub[18] = f"Results after R{k} for each table" if k == 1 else f"Rearranged results after R{k}"
        if kind == "r1":
            sub[23] = f"Scores after R{k}, alphabetically"; sub[28] = f"Scores after R{k}, ranked"
        elif kind == "mid":
            sub[23] = f"Scores after R{k}, alphabetically"; sub[28] = f"Scores after R{k}, ranked"
        elif kind == "last":
            sub[23] = f"Scores after R{k}, by rank"
        else:
            sub[23] = "Final standings"
        for c, v in sub.items():
            _put(ws, title_row + 1, c, v)

        prev_t = "TurnPR1" if k == 2 else f"ZTurnPR{k-1}"
        prev_r = "RelScR1" if k == 2 else f"ZRelScR{k-1}"
        if kind == "r1":
            hdr = {1: "Table", 2: "StartNo", 3: "Name", 6: f"ScoreR{k}", 7: f"RelScR{k}", 8: f"TabRkR{k}", 9: f"TurnPR{k}",
                   10: f"TabRkR{k}", 11: f"TurnPR{k}", 18: f"TabRkR{k}", 19: "Name", 20: f"TurnPR{k}", 21: f"RelScR{k}",
                   23: f"TabRankR{k}", 24: "Name", 25: f"ZTurnPR{k}", 26: f"ZRelScR{k}",
                   28: f"RankR{k}", 29: "Name", 30: f"ZTurnPR{k}", 31: f"ZRelScR{k}"}
        else:
            hdr = {1: "Table", 2: "StartNo", 3: "Name", 4: prev_t, 5: prev_r, 6: f"ScoreR{k}", 7: f"RelScR{k}",
                   8: f"TabRkR{k}", 9: f"TurnPR{k}p", 10: f"TabRkR{k}", 11: f"TurnPR{k}",
                   13: f"ZTurnPR{k}", 14: f"ZRelScR{k}", 15: f"RankR{k}p", 16: f"RankR{k}",
                   18: f"RankR{k}", 19: "Name", 20: f"ZTurnPR{k}", 21: f"ZRelScR{k}",
                   23: f"RankR{k}", 24: "Name", 25: f"ZTurnPR{k}", 26: f"ZRelScR{k}"}
            if kind == "mid":
                hdr.update({28: f"RankR{k}", 29: "Name", 30: f"ZTurnPR{k}", 31: f"ZRelScR{k}"})
            if kind == "final":
                hdr.pop(2)
        for c, v in hdr.items():
            _put(ws, title_row + 2, c, v)

        # ── data rows ──────────────────────────────────────────────────────
        order = [x for _, keys, _ in groups for x in keys]
        for x in order:
            r = row_of[x]
            gs, ge = group_range[x]
            src = (sty[7] if r == de else sty[3] if r == ds else
                   sty[4] if r == gs else sty[6] if r == ge else sty[5])
            st.row(src, r)
            grp = next(g for g in groups if x in g[1])
            is_abs = grp[2]
            _put(ws, r, 1, "" if is_abs else grp[0])
            if kind != "final":
                _put(ws, r, 2, start_no[x])
                _put(ws, r, 3, f"=LOOKUP(B{r},{LIST_H},{LIST_F})")
            if k > 1:
                if kind == "final":
                    pr = prev["rank_row"][x]
                    pc = prev["rank_cols"]
                    _put(ws, r, 3, f"={pc[0]}{pr}")
                    _put(ws, r, 4, f"={pc[1]}{pr}")
                    _put(ws, r, 5, f"={pc[2]}{pr}")
                else:
                    a0, a1, ac = prev["alpha"]
                    _put(ws, r, 4, f"=INDEX({ac[1]}${a0}:{ac[1]}${a1},MATCH(C{r},{ac[0]}${a0}:{ac[0]}${a1},0))")
                    _put(ws, r, 5, f"=INDEX({ac[2]}${a0}:{ac[2]}${a1},MATCH(C{r},{ac[0]}${a0}:{ac[0]}${a1},0))")
            if is_abs:
                _put(ws, r, 7, 0, "Did not play this round – no score registered.")
                _put(ws, r, 11, 0)
            else:
                _put(ws, r, 6, F[x])
                _put(ws, r, 7, f"=F{r}-AVERAGE(F${gs}:F${ge})")
                _put(ws, r, 8, f"=RANK(F{r},F${gs}:F${ge})")
                pts = f"IF(H{r}=1,10,IF(H{r}=2,6,IF(H{r}=3,3,1)))"
                _put(ws, r, 9, f"={weight_ref}*{pts}" if kind == "final" else f"={pts}")
                _put(ws, r, 10, H[x], TIE_NOTE if x in ties else None)
                _put(ws, r, 11, K[x])
            if k == 1:
                _put(ws, r, 18, f"=J{r}"); _put(ws, r, 19, f"=C{r}"); _put(ws, r, 20, f"=K{r}"); _put(ws, r, 21, f"=G{r}")
            else:
                _put(ws, r, 13, f"=D{r}+K{r}")
                _put(ws, r, 14, f"=E{r}+G{r}")
                if kind == "final":
                    off = gs - ds
                    rank_f = f"RANK(M{r},M${gs}:M${ge},0)"
                    _put(ws, r, 15, f"={off}+{rank_f}" if off else f"={rank_f}")
                else:
                    _put(ws, r, 15, f"=RANK(M{r},M${ds}:M${de},0)")
                _put(ws, r, 16, Pr[x], "Tied on both ZTurnP and ZRelSc – ordered alphabetically." if x in full_ties else None)
                _put(ws, r, 18, f"=P{r}"); _put(ws, r, 19, f"=C{r}"); _put(ws, r, 20, f"=M{r}"); _put(ws, r, 21, f"=N{r}")

        # W..Z and AB..AE lists (pasted values in the template)
        if kind in ("r1", "mid"):
            alpha = sorted(P, key=lambda x: _nb_key(P[x]["name"]))
            for i, x in enumerate(alpha):
                r = ds + i
                _put(ws, r, 23, H[x] if k == 1 else Pr[x])
                _put(ws, r, 24, P[x]["name"])
                _put(ws, r, 25, M[x]); _put(ws, r, 26, Nn[x])
            ranked = sorted(P, key=lambda x: Pr[x])
            for i, x in enumerate(ranked):
                r = ds + i
                _put(ws, r, 28, i + 1); _put(ws, r, 29, P[x]["name"]); _put(ws, r, 30, M[x]); _put(ws, r, 31, Nn[x])
            rank_row = {x: ds + i for i, x in enumerate(ranked)}
            prev = {"alpha": (ds, de, ("X", "Y", "Z")), "rank": Pr, "rank_row": rank_row, "rank_cols": ("AC", "AD", "AE")}
            standings_src = ("ab", "ae", ds, de, [(i + 1, P[x]["name"], M[x], Nn[x]) for i, x in enumerate(ranked)])
        else:
            ranked = sorted(P, key=lambda x: Pr[x])
            for i, x in enumerate(ranked):
                r = ds + i
                _put(ws, r, 23, Pr[x]); _put(ws, r, 24, P[x]["name"]); _put(ws, r, 25, M[x]); _put(ws, r, 26, Nn[x])
            rank_row = {x: ds + i for i, x in enumerate(ranked)}
            prev = {"alpha": None, "rank": Pr, "rank_row": rank_row, "rank_cols": ("X", "Y", "Z")}
            standings_src = ("w", "z", ds, de, [(Pr[x], P[x]["name"], M[x], Nn[x]) for x in ranked])

        cum = {x: (M[x], Nn[x]) for x in P}
        title_row = de + 3

    # ── Final standings ────────────────────────────────────────────────────
    fs = FINAL_STANDINGS_STYLE
    st.cell(fs["title"], 1, title_row); st.cell(fs["title"], 2, title_row)
    _put(ws, title_row, 1, "FINAL STANDINGS")
    st.cell(fs["note"], 1, title_row + 1)
    c0, c1, s0, s1, rows = standings_src
    _put(ws, title_row + 1, 1, f"Copy {c0}{s0}:{c1}{s1}, paste values")
    for c in range(1, 5):
        st.cell(fs["header"], c, title_row + 2)
    for c, h in enumerate(("Rank", "Name", "TurnP", "RelSc"), start=1):
        _put(ws, title_row + 2, c, h)
    group_ends, acc = set(), 0
    for size in (final_group_sizes or [4] * ((len(rows) + 3) // 4)):
        acc += size
        group_ends.add(acc - 1)
    for i, (rank, name, tp, rs) in enumerate(rows):
        r = title_row + 3 + i
        src = (fs["first"] if i == 0 else fs["last"] if i == len(rows) - 1 else
               fs["group_last"] if i in group_ends else fs["mid"])
        for c in range(1, 5):
            st.cell(src, c, r)
        _put(ws, r, 1, rank); _put(ws, r, 2, name); _put(ws, r, 3, tp); _put(ws, r, 4, rs)

    if notes:
        _put(ws, 3, 11, "Notes: " + " ".join(notes))

    # sheet-level formatting
    for col, dim in src_o.column_dimensions.items():
        ws.column_dimensions[col].width = dim.width
    ws.sheet_format = copy(src_o.sheet_format)
    ws.sheet_view.zoomScale = src_o.sheet_view.zoomScale

    # ── Printout (table overview per game, final excluded as in the template) ──
    wp = wb.create_sheet("_Printout")
    sp = _Styler(src_p, wp)
    games = rounds[:n_normal]
    ncols = 1 + len(games)
    for c in range(1, ncols + 1):
        sp.cell(1, min(c, 6), 1, c)
    _put(wp, 1, 1, "Bord")
    for gi, rnd in enumerate(games):
        _put(wp, 1, 2 + gi, f"Spill {gi + 1}")
    seats: dict = {}
    for gi, rnd in enumerate(games):
        for t, keys in rnd["tables"]:
            seats.setdefault(t, {})[gi] = keys
    seat_pos = {}
    for s in scores:
        g = _game_no(s.get("gameNumber"))
        if g is not None:
            seat_pos[(g, _ident(s.get("name") or ""))] = _to_int(s.get("startingPosition"))
    r = 2
    tables_sorted = sorted(seats, key=lambda t: (t == "?", t if t != "?" else 0))
    for ti, t in enumerate(tables_sorted):
        height = max(len(v) for v in seats[t].values())
        base = (2, 3, 5) if ti % 2 == 0 else (6, 7, 9)
        for i in range(height):
            src = base[0] if i == 0 else base[2] if i == height - 1 else base[1]
            for c in range(1, ncols + 1):
                sp.cell(src, min(c, 6), r + i, c)
            _put(wp, r + i, 1, t)
        for gi, keys in seats[t].items():
            g = games[gi]["game"]
            ordered = sorted(keys, key=lambda x: (seat_pos.get((g, x)) or 99, start_no.get(x, 999)))
            for i, x in enumerate(ordered):
                _put(wp, r + i, 2 + gi, P[x]["name"] if x in P else x)
        r += height
    for col, dim in src_p.column_dimensions.items():
        wp.column_dimensions[col].width = dim.width
    for c in range(7, ncols + 1):
        wp.column_dimensions[L(c)].width = 11.0

    # swap in the generated sheets
    wb.remove(src_o); wb.remove(src_p)
    ws.title, wp.title = "Oppsett", "Printout"
    wb.active = 0

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
